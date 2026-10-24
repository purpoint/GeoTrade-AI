"""Portfolio file import (P2-BE-04). File parsing only: Zerodha (Kite CSV, Console XLSX), Groww, generic.
Rows come back as a preview; nothing is saved until the user confirms."""

import csv
import io
import re
import zipfile
from dataclasses import asdict, dataclass
from functools import lru_cache
from typing import Any

from geotrade_ai import reference

MAX_BYTES = 1024 * 1024
MAX_ROWS = 500
# symbols renamed or demerged since brokers exported them
RENAMES = {"ZOMATO": "ETERNAL", "TMPV": "TATAMOTORS", "MM": "M&M", "BAJAJAUTO": "BAJAJ-AUTO"}
SERIES_SUFFIX = re.compile(r"(-(EQ|BE|BZ|BL|SM|ST)|\.(NS|BO|NSE|BSE))$", re.I)
FORMULA_START = ("=", "+", "-", "@", "\t", "\r")
NAME_NOISE = re.compile(r"\b(ltd|limited|ltd\.|co|company|corporation|corp|india|the)\b\.?", re.I)

FORMATS: dict[str, dict[str, tuple[str, ...]]] = {
    "zerodha_kite": {"symbol": ("instrument",), "quantity": ("qty.",), "avg_price": ("avg. cost",)},
    "zerodha_console": {
        "symbol": ("symbol",), "isin": ("isin",), "quantity": ("quantity available", "quantity"),
        "avg_price": ("average price",),
    },
    "groww": {"name": ("stock name", "company name"), "isin": ("isin",), "quantity": ("quantity",),
              "avg_price": ("average buy price", "average price", "avg. price")},
}  # fmt: skip
GENERIC_GUESS = {
    "symbol": ("symbol", "ticker", "instrument", "stock", "scrip", "trading symbol", "name"),
    "quantity": ("quantity", "qty", "qty.", "shares", "units"),
    "avg_price": (
        "avg price",
        "average price",
        "avg. cost",
        "avg cost",
        "buy price",
        "price paid",
        "cost",
        "price",
    ),
    "isin": ("isin",),
}


class ImportError_(ValueError):
    """User-facing import problem (bad file, too big, unknown layout)."""


@dataclass
class ParsedRow:
    symbol: str
    ticker: str | None
    quantity: float | None
    avg_price: float | None
    isin: str | None = None
    status: str = "matched"  # matched | unmatched | invalid
    issue: str | None = None

    def as_dict(self) -> dict[str, Any]:
        return asdict(self)


def neutralise(value: str) -> str:
    """Stop spreadsheet formula injection when the value is shown or exported again."""
    value = value.strip()
    return "'" + value if value.startswith(FORMULA_START) else value


def _number(value: Any) -> float | None:
    if value is None:
        return None
    if isinstance(value, int | float):
        return float(value)
    text = str(value).strip().replace(",", "").replace("₹", "")
    if not text or text.startswith(("=", "@")):
        return None
    try:
        return float(text)
    except ValueError:
        return None


# ---------- reading files ----------


def read_table(filename: str, content: bytes) -> list[list[str]]:
    if len(content) > MAX_BYTES:
        raise ImportError_("file is larger than 1 MB")
    name = filename.lower()
    if name.endswith(".xlsx"):
        return _read_xlsx(content)
    if name.endswith((".csv", ".txt")):
        return _read_csv(content)
    raise ImportError_("upload a .csv or .xlsx file (.xlsm and .xls are not accepted)")


def _read_csv(content: bytes) -> list[list[str]]:
    try:
        text = content.decode("utf-8-sig")
    except UnicodeDecodeError:
        text = content.decode("latin-1")
    try:
        dialect: Any = csv.Sniffer().sniff(text[:4096], delimiters=",;\t")
    except csv.Error:
        dialect = csv.excel
    return [row for row in csv.reader(io.StringIO(text), dialect)]


def _read_xlsx(content: bytes) -> list[list[str]]:
    from openpyxl import load_workbook

    if not zipfile.is_zipfile(io.BytesIO(content)):
        raise ImportError_("this .xlsx file is not a valid workbook")
    with zipfile.ZipFile(io.BytesIO(content)) as z:
        if any(n.lower().endswith("vbaproject.bin") for n in z.namelist()):
            raise ImportError_("workbooks with macros are not accepted")
    # read-only and values only: formulas are never evaluated, macros never run
    wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    try:
        ws = wb.worksheets[0]
        return [
            ["" if c is None else str(c) for c in row]
            for row in ws.iter_rows(values_only=True, max_row=MAX_ROWS + 50)
        ]
    finally:
        wb.close()


# ---------- format detection ----------


def _norm(h: str) -> str:
    return re.sub(r"\s+", " ", h.strip().lower())


def _find_header(rows: list[list[str]], extra: set[str] | None = None) -> tuple[int, list[str]]:
    """Brokers put title rows above the table; the header is the first row naming a known column."""
    known = (
        {n for f in FORMATS.values() for names in f.values() for n in names}
        | {n for names in GENERIC_GUESS.values() for n in names}
        | (extra or set())
    )
    for i, row in enumerate(rows[:30]):
        cells = [_norm(c) for c in row]
        if sum(c in known for c in cells) >= 2:
            return i, cells
    raise ImportError_("could not find a header row with symbol and quantity columns")


def detect_format(headers: list[str]) -> str:
    for name, spec in FORMATS.items():
        if all(any(alias in headers for alias in aliases) for aliases in spec.values() if aliases):
            return name
    return "generic"


def _column(headers: list[str], aliases: tuple[str, ...]) -> int | None:
    for alias in aliases:
        if alias in headers:
            return headers.index(alias)
    return None


# ---------- symbol matching ----------


@lru_cache
def _symbol_index() -> tuple[dict[str, str], dict[str, str]]:
    tickers = {c["ticker"]: c["ticker"] for c in reference.companies()}
    names: dict[str, str] = {}
    for c in reference.companies():
        names[_clean_name(c["name"])] = c["ticker"]
    for a in reference.aliases():
        if not a["is_ambiguous"]:
            names.setdefault(_clean_name(str(a["alias"])), str(a["ticker"]))
    return tickers, names


def _clean_name(name: str) -> str:
    return re.sub(r"[^a-z0-9&]+", " ", NAME_NOISE.sub(" ", name.lower())).strip()


def normalise_symbol(raw: str) -> str:
    s = SERIES_SUFFIX.sub("", raw.strip().upper())
    return RENAMES.get(s, s)


def match_ticker(symbol: str | None, name: str | None = None) -> str | None:
    tickers, names = _symbol_index()
    if symbol:
        s = normalise_symbol(symbol)
        if s in tickers:
            return s
    if name or symbol:
        return names.get(_clean_name(name or symbol or ""))
    return None


# ---------- parse ----------


def parse(
    filename: str, content: bytes, mapping: dict[str, str] | None = None
) -> tuple[str, list[ParsedRow]]:
    rows = read_table(filename, content)
    start, headers = _find_header(rows, {_norm(v) for v in (mapping or {}).values() if v})
    fmt = detect_format(headers)
    spec: dict[str, tuple[str, ...]] = dict(FORMATS.get(fmt, GENERIC_GUESS))
    if fmt == "generic" and mapping:
        spec.update({k: (_norm(v),) for k, v in mapping.items() if k in GENERIC_GUESS and v})
    cols = {key: _column(headers, aliases) for key, aliases in spec.items()}
    if cols.get("quantity") is None or (cols.get("symbol") is None and cols.get("name") is None):
        raise ImportError_("map the symbol and quantity columns to import this file")

    def cell(row: list[str], key: str) -> str:
        i = cols.get(key)
        return row[i].strip() if i is not None and i < len(row) else ""

    body = [r for r in rows[start + 1 :] if any(c.strip() for c in r)]
    if len(body) > MAX_ROWS:
        raise ImportError_(f"file has more than {MAX_ROWS} rows")
    parsed = []
    for r in body:
        raw_symbol, raw_name = cell(r, "symbol"), cell(r, "name")
        if not (raw_symbol or raw_name) or _norm(raw_symbol or raw_name) in ("total", "grand total"):
            continue
        shown = neutralise(raw_symbol or raw_name)
        injected = shown.startswith("'")
        ticker = None if injected else match_ticker(raw_symbol or None, raw_name or None)
        # matched rows are keyed by ticker; unmatched keep the broker's symbol (or the name, for Groww)
        symbol = ticker or (normalise_symbol(shown) if raw_symbol and not injected else shown)
        qty, price = _number(cell(r, "quantity")), _number(cell(r, "avg_price"))
        row = ParsedRow(
            symbol=symbol,
            ticker=ticker,
            quantity=qty,
            avg_price=price,
            isin=neutralise(cell(r, "isin")) or None,
        )
        if injected:
            row.status, row.issue = "invalid", "looks like a spreadsheet formula"
        elif qty is None or qty <= 0:
            row.status, row.issue = "invalid", "quantity must be a positive number"
        elif price is None or price < 0:
            row.status, row.issue = "invalid", "missing average price"
        elif ticker is None:
            row.status, row.issue = "unmatched", "not a NIFTY 50 stock; kept without news signals"
        parsed.append(row)
    return fmt, merge_duplicates(parsed)


def merge_duplicates(rows: list[ParsedRow]) -> list[ParsedRow]:
    """Same stock twice (lots, exchanges): add quantities, weight the average price."""
    out: dict[str, ParsedRow] = {}
    for row in rows:
        key = row.ticker or row.symbol
        prev = out.get(key)
        if prev is None or row.status == "invalid" or prev.status == "invalid":
            out[key if prev is None else f"{key}#{len(out)}"] = row
            continue
        assert prev.quantity is not None and row.quantity is not None
        qty = prev.quantity + row.quantity
        prev.avg_price = round(
            ((prev.avg_price or 0) * prev.quantity + (row.avg_price or 0) * row.quantity) / qty, 4
        )
        prev.quantity = qty
    return list(out.values())


SAMPLE_PORTFOLIO = [
    ("RELIANCE", 25, 2480.0), ("HDFCBANK", 40, 1520.0), ("INFY", 35, 1480.0), ("TCS", 10, 3650.0),
    ("ICICIBANK", 30, 1050.0), ("ITC", 150, 430.0), ("LT", 8, 3400.0), ("BHARTIARTL", 15, 1350.0),
    ("SUNPHARMA", 20, 1600.0), ("TATAMOTORS", 60, 720.0), ("ADANIENT", 10, 2600.0), ("MARUTI", 3, 11800.0),
]  # fmt: skip
